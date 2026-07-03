"""Выгрузка результатов в CSV и Excel."""
from __future__ import annotations

import csv
import os
from typing import Iterable

from .models import Company

COLUMNS = [
    "ИНН", "ОГРН", "Название", "Полное название", "Основной ОКВЭД",
    "Вид деятельности", "Регион", "Статус", "Адрес",
    "Телефоны", "Почты", "Сайты", "Источник контактов", "Ошибка обогащения",
]


def read_existing_keys(path: str) -> set[str]:
    """ИНН/ОГРН уже собранных компаний из CSV (для докачки)."""
    keys: set[str] = set()
    if os.path.exists(path):
        with open(path, encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f, delimiter=";"):
                k = (row.get("ИНН") or "").strip() or (row.get("ОГРН") or "").strip()
                if k:
                    keys.add(k)
    return keys


class CsvAppender:
    """Потоковая дозапись компаний в CSV (не теряется при остановке)."""

    def __init__(self, path: str):
        exists = os.path.exists(path) and os.path.getsize(path) > 0
        self.path = path
        self.f = open(path, "a", encoding="utf-8-sig", newline="")
        self.w = csv.DictWriter(self.f, fieldnames=COLUMNS, delimiter=";")
        if not exists:
            self.w.writeheader()
            self.f.flush()

    def write(self, company: Company) -> None:
        self.w.writerow(company.to_row())
        self.f.flush()

    def close(self) -> None:
        try:
            self.f.close()
        except Exception:  # noqa: BLE001
            pass


def write_excel_from_csv(csv_path: str, xlsx_path: str) -> int:
    """Строит Excel из готового CSV (чтобы включить и дозаписанные строки)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Компании"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    n = 0
    if os.path.exists(csv_path):
        with open(csv_path, encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f, delimiter=";"):
                ws.append([row.get(col, "") for col in COLUMNS])
                n += 1
    widths = {"Название": 40, "Полное название": 50, "Вид деятельности": 45,
              "Адрес": 50, "Телефоны": 25, "Почты": 30, "Сайты": 30}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 16)
    wb.save(xlsx_path)
    return n


def write_csv(companies: Iterable[Company], path: str) -> int:
    n = 0
    # utf-8-sig — чтобы кириллица корректно открывалась в Excel
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, delimiter=";")
        writer.writeheader()
        for c in companies:
            writer.writerow(c.to_row())
            n += 1
    return n


def write_excel(companies: Iterable[Company], path: str) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Компании"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    n = 0
    for c in companies:
        row = c.to_row()
        ws.append([row[col] for col in COLUMNS])
        n += 1

    # ширина колонок (грубая автоподгонка)
    widths = {"Название": 40, "Полное название": 50, "Вид деятельности": 45,
              "Адрес": 50, "Телефоны": 25, "Почты": 30, "Сайты": 30}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 16)

    wb.save(path)
    return n
